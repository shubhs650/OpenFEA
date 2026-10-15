"""Real legacy-model and long-table DOCX parity, conversion, layout and editability."""

import hashlib
import io
import json
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.domain.reports import ReportModel, ReportTable
from tests.reports.test_reports import assert_a4_bounds, pdf_text
from tests.reports.word_support import convert_word
from worker.reports.content import ATTRIBUTION, assumption, rows
from worker.reports.docx import EDIT_NOTICE, validate_docx
from worker.reports.html import report_html
from worker.reports.render import render_pdf, render_word, word_generator

LEGACY = Path("docs/validation/results/T025/historical.json")
OUT = Path(".local/reports/T026")


def parity(model: ReportModel, raw: bytes) -> None:
    validate_docx(raw)
    document = Document(io.BytesIO(raw))
    paragraphs = [p.text for p in document.paragraphs]
    assert model.request.title in paragraphs
    assert ATTRIBUTION in paragraphs and EDIT_NOTICE in paragraphs
    assert (model.request.user_conclusion or "No user conclusion supplied.") in paragraphs
    assert all(w in paragraphs for w in model.warnings)
    assert all(assumption(n) in paragraphs for n in model.assumptions)
    assert all(f"Figure {i + 1}. {f.caption}" in paragraphs for i, f in enumerate(model.figures))
    assert len(document.tables) == len(model.tables)
    for actual, expected in zip(document.tables, model.tables, strict=True):
        assert [cell.text for cell in actual.rows[0].cells] == expected.columns
        assert [tuple(cell.text for cell in row.cells) for row in actual.rows[1:]] == list(
            rows(expected)
        )
        assert actual.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert len(document.inline_shapes) == len(model.figures)
    for image in document.inline_shapes:
        assert int(image.width) == 6480000  # 180 mm, available A4 width.
        assert image.height is not None and 0 < int(image.height) <= 6480000
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        assert all(not p.filename.endswith(".pdf") for p in archive.infolist())
        assert len([p for p in archive.namelist() if p.startswith("word/media/")]) == len(
            model.figures
        )


@pytest.mark.parametrize("kind", ["legacy", "long-special"])
def test_actual_word_parity_editing_and_converted_pages(kind: str, tmp_path: Path) -> None:
    model = ReportModel.model_validate_json(LEGACY.read_bytes())
    if kind == "long-special":
        request = model.request.model_copy(
            update={
                "title": "Évaluation σxx ≤ 1.2 MPa & <review> μm",
                "user_conclusion": "Conclusion écrite: σ = 1.2 MPa; μm; <review> & author.",
            }
        )
        table = ReportTable(
            title="Long material table — σ & μm",
            columns=["Material provenance row", "Source / condition"],
            rows=[
                [
                    f"Material {i:03d}",
                    "Évaluation σxx ≤ 1.2 MPa & <review> μm " * 6 + f"END-{i:03d}",
                ]
                for i in range(90)
            ],
        )
        model = model.model_copy(update={"request": request, "tables": [table, *model.tables]})
    word = tmp_path / "word"
    word.mkdir()
    pdf = tmp_path / "pdf"
    pdf.mkdir()
    if kind == "legacy":
        (word / "model.json").write_bytes(LEGACY.read_bytes())
        shutil.copyfile(LEGACY.with_suffix(".pdf"), pdf / "report.pdf")
    else:
        (word / "model.json").write_text(model.model_dump_json())
        (pdf / "report.html").write_text(report_html(model))
        render_pdf(pdf)
    model_sha = hashlib.sha256((word / "model.json").read_bytes()).hexdigest()
    provenance = {
        "source_model_sha256": model_sha,
        "original_pdf_sha256": hashlib.sha256((pdf / "report.pdf").read_bytes()).hexdigest(),
        **word_generator(),
    }
    (word / "provenance.json").write_text(json.dumps(provenance))
    audit = render_word(word)
    raw = (word / "report.docx").read_bytes()
    parity(model, raw)
    # Paragraphs AND tables can be edited as native Word content and reopened.
    editable = Document(io.BytesIO(raw))
    editable.paragraphs[0].text = "User edited this title"
    editable.tables[0].cell(1, 0).text = "User edited this cell"
    edited_path = tmp_path / "edited.docx"
    editable.save(str(edited_path))
    reopened = Document(str(edited_path))
    assert reopened.paragraphs[0].text == "User edited this title"
    assert reopened.tables[0].cell(1, 0).text == "User edited this cell"
    assert edited_path.read_bytes() != raw and (word / "report.docx").read_bytes() == raw
    converted = convert_word(word)
    text = pdf_text(converted)
    assert "User-written conclusion" in text and "Word export provenance" in text
    assert str(model.source.job_id) in text
    pages = assert_a4_bounds(converted)
    assert_a4_bounds(pdf / "report.pdf")
    if kind == "long-special":
        assert "END-000" in text and "END-089" in text
        assert text.count("Material provenance row") >= 3
        assert "Évaluation" in text and "<review> & author." in text
    OUT.mkdir(parents=True, exist_ok=True)
    for source, name in [
        (word / "model.json", f"{kind}-model.json"),
        (word / "report.docx", f"{kind}.docx"),
        (converted, f"{kind}-word.pdf"),
        (pdf / "report.pdf", f"{kind}-direct.pdf"),
    ]:
        shutil.copyfile(source, OUT / name)
    (OUT / f"{kind}-audit.json").write_text(
        json.dumps(
            {
                "word_pages": pages,
                "renderer": audit,
                "generator": provenance,
                "docx_sha256": hashlib.sha256(raw).hexdigest(),
            },
            indent=2,
        )
        + "\n"
    )
