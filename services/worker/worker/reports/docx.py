"""Editable paragraphs/tables with bounded embedded raster copies of the shared figures."""

import io
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt, RGBColor

from app.domain.reports import ReportModel
from worker.reports.content import (
    ATTRIBUTION,
    CLOSING,
    assumption,
    identity,
    overview,
    reproducibility,
    rows,
)

EDIT_NOTICE = (
    "Editable Word copy. The recorded output checksum and run provenance identify the generated "
    "version only; they do not certify later user edits."
)


def valid_text(text: str) -> None:
    if any(
        not (
            c in "\t\n\r"
            or 0x20 <= ord(c) <= 0xD7FF
            or 0xE000 <= ord(c) <= 0xFFFD
            or 0x10000 <= ord(c) <= 0x10FFFF
        )
        for c in text
    ):
        raise ValueError("Report text contains a character unsupported by editable Word XML")


def validate_docx(raw: bytes) -> None:
    if len(raw) > 16 * 1024**2:
        raise ValueError("Word output exceeds 16 MiB")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        files = archive.infolist()
        if len(files) > 64 or sum(f.file_size for f in files) > 32 * 1024**2:
            raise ValueError("Word output exceeds package bounds")
        if "word/document.xml" not in archive.namelist():
            raise ValueError("Missing editable Word document")
        for entry in files:
            if entry.filename.startswith("/") or ".." in entry.filename.split("/"):
                raise ValueError("Invalid Word package path")
            if "vba" in entry.filename.lower() or entry.flag_bits & 1:
                raise ValueError("Executable/encrypted Word output is forbidden")
            if entry.filename.endswith(".rels"):
                from xml.etree import ElementTree as ET

                root = ET.fromstring(archive.read(entry))
                if any(r.attrib.get("TargetMode") == "External" for r in root):
                    raise ValueError("External Word relationships are forbidden")


def write_docx(model: ReportModel, directory: Path, provenance: dict[str, str]) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(17), Mm(19)
    section.left_margin = section.right_margin = Mm(15)
    section.footer_distance = Mm(9)
    for name, size in (("Normal", 10), ("Title", 23), ("Heading 1", 14)):
        style = document.styles[name]
        style.font.name, style.font.size = "DejaVu Sans", Pt(size)
        style.font.color.rgb = RGBColor.from_string("172633")
    normal = document.styles["Normal"].paragraph_format
    normal.space_after, normal.line_spacing = Pt(6), 1.15
    document.core_properties.title = model.request.title
    document.core_properties.author = "OpenFEA report generator"
    document.core_properties.subject = f"Frozen run {model.source.job_id}"
    document.core_properties.identifier = f"report:{model.id}"
    document.core_properties.comments = (
        "Source model SHA-256 "
        + provenance.get("source_model_sha256", "Unavailable")
        + "; PDF SHA-256 "
        + provenance.get("original_pdf_sha256", "Unavailable")
    )
    document.core_properties.created = document.core_properties.modified = model.created_at

    def paragraph(text: str, *, small: bool = False, heading: bool = False) -> Any:
        valid_text(text)
        p = document.add_paragraph(text, style="Heading 1" if heading else None)
        if small:
            for run in p.runs:
                run.font.size = Pt(8)
        return p

    def cell_text(cell: Any, text: str, *, header: bool = False) -> None:
        valid_text(text)
        cell.text = text
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.bold = header

    valid_text(model.request.title)
    document.add_heading(model.request.title, 0)
    paragraph(overview(model))
    paragraph(identity(model), small=True)
    paragraph("Scope, warnings and assumptions", heading=True)
    for warning in model.warnings:
        p = paragraph(warning)
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("894700")
    for note in model.assumptions:
        paragraph(assumption(note))
    for limitation in model.limitations:
        paragraph(limitation)
    paragraph("Selected figures and setup", heading=True)
    for index, figure in enumerate(model.figures):
        image = directory / f"figure-{index}.png"
        if image.stat().st_size > 8 * 1024**2 or not image.read_bytes().startswith(
            b"\x89PNG\r\n\x1a\n"
        ):
            raise ValueError("Word figure is unavailable or exceeds its 8 MiB bound")
        p = document.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(image), width=Mm(180))
        caption = paragraph(f"Figure {index + 1}. {figure.caption}")
        caption.paragraph_format.keep_together = True
        for run in caption.runs:
            run.font.size = Pt(9)
    for table in model.tables:
        paragraph(table.title, heading=True)
        target = document.add_table(rows=1, cols=len(table.columns))
        target.style = "Table Grid"
        target.autofit = False
        for column in target.columns:
            column.width = Mm(180 / len(table.columns))
        for cell, text in zip(target.rows[0].cells, table.columns, strict=True):
            cell_text(cell, text, header=True)
        properties = target.rows[0]._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:tblHeader"))
        for row in rows(table):
            target_row = target.add_row()
            target_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, text in zip(target_row.cells, row, strict=True):
                cell_text(cell, text)
        document.add_paragraph()
    paragraph("User-written conclusion", heading=True)
    paragraph(ATTRIBUTION)
    paragraph(model.request.user_conclusion or "No user conclusion supplied.")
    paragraph("Reproducibility IDs", heading=True)
    for key, v in reproducibility(model).items():
        paragraph(f"{key}: {v}", small=True)
    paragraph(CLOSING)
    paragraph("Word export provenance", heading=True)
    paragraph(EDIT_NOTICE)
    for key, v in provenance.items():
        paragraph(f"{key}: {v}", small=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    run = footer.add_run("Engineering report · ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr", "PAGE")
    footer._p.append(field)
    output = directory / "report.docx"
    document.save(str(output))
    validate_docx(output.read_bytes())


def validate_text_tree(value: Any) -> None:
    if isinstance(value, str):
        valid_text(value)
    elif isinstance(value, list):
        for item in value:
            validate_text_tree(item)
    elif isinstance(value, dict):
        for item in value.values():
            validate_text_tree(item)
